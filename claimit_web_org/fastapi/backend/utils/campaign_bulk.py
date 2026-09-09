"""
campaign_bulk.py — read the WhatsApp campaign sheet an admin uploads.

The sheet is filled in by hand from bulk-upload data, so it arrives messy:
numbers stored as Excel numbers (so 9842512345 becomes 9842512345.0), leading
apostrophes, +91 prefixes, spaces and dashes, blank rows in the middle, the
grey instruction row still present, and the same shop listed twice.

Every one of those is handled here rather than at send time, because a bad row
that reaches the sender costs money or messages the wrong person.

Only two columns matter:
    phone       REQUIRED — the only field a row cannot do without
    shop_name   optional — fills {{1}}; blank falls back to a neutral phrase
city and pincode may be present for the admin's own filtering; they are read
but never sent.
"""
import io
import re
from typing import List, Tuple

from openpyxl import load_workbook

REQUIRED = ("phone",)
OPTIONAL = ("shop_name", "city", "pincode")

# The note row the template ships with. If admin leaves it in, skip it rather
# than trying to message "REQUIRED. WhatsApp number...".
_NOTE_MARKERS = ("required.", "optional.", "not sent")


def _clean_phone(value) -> str:
    """Any of the shapes Excel produces -> a bare 10-digit Indian mobile.

    Returns "" when the value can't be a mobile number, which the caller
    reports as a rejected row instead of trying to send to it.
    """
    if value is None:
        return ""
    # Excel stores a typed number as a float: 9842512345 -> 9842512345.0
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)

    text = text.strip().lstrip("'")               # leading apostrophe
    digits = re.sub(r"\D", "", text)              # drop +, spaces, dashes, ()

    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]                       # +91 form
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]                       # STD 0 prefix

    if len(digits) != 10:
        return ""
    if digits[0] not in "6789":                   # Indian mobiles start 6-9
        return ""
    return digits


def _clean_text(value) -> str:
    s = str(value or "").strip()
    return "" if s.lower() in ("none", "null", "nan", "-") else s


def parse_campaign_sheet(data: bytes) -> Tuple[List[dict], List[dict], dict]:
    """(recipients, rejected, summary).

    `recipients` is deduplicated and ready to send. `rejected` explains every
    row that was dropped and why, so the admin can fix the sheet rather than
    wonder why 40 shops never heard from them.
    """
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["Send List"] if "Send List" in wb.sheetnames else wb.worksheets[0]

    # Header row: first row containing "phone".
    header_row, headers = None, {}
    for r in range(1, min(ws.max_row, 10) + 1):
        values = [str(ws.cell(row=r, column=c).value or "").strip().lower()
                  for c in range(1, ws.max_column + 1)]
        if "phone" in values:
            header_row = r
            headers = {name: i + 1 for i, name in enumerate(values) if name}
            break
    if header_row is None:
        raise ValueError(
            "Could not find a 'phone' column. Use the downloaded template and "
            "do not rename the headers.")

    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")

    recipients, rejected, seen = [], [], set()

    for r in range(header_row + 1, ws.max_row + 1):
        raw_phone = ws.cell(row=r, column=headers["phone"]).value
        raw_name = (ws.cell(row=r, column=headers["shop_name"]).value
                    if "shop_name" in headers else None)

        # Skip the instruction row and fully blank rows without reporting them
        # as errors — neither is a mistake by the person filling the sheet.
        joined = f"{raw_phone} {raw_name}".lower()
        if any(m in joined for m in _NOTE_MARKERS):
            continue
        if raw_phone is None and raw_name is None:
            continue

        name = _clean_text(raw_name)
        phone = _clean_phone(raw_phone)

        # The mobile number is the only thing that MUST be there — without it
        # there is nobody to message. A missing shop name is not a reason to
        # skip a real number; the sender fills {{1}} with a neutral phrase.
        if not phone:
            rejected.append({"row": r, "phone": _clean_text(raw_phone),
                             "shop_name": name,
                             "reason": "Not a valid 10-digit mobile number"})
            continue
        if phone in seen:
            rejected.append({"row": r, "phone": phone, "shop_name": name,
                             "reason": "Duplicate number in this sheet"})
            continue

        seen.add(phone)
        recipients.append({
            "phone": phone,
            "shop_name": name,
            "city": _clean_text(ws.cell(row=r, column=headers["city"]).value)
            if "city" in headers else "",
            "pincode": _clean_text(ws.cell(row=r, column=headers["pincode"]).value)
            if "pincode" in headers else "",
        })

    summary = {
        "rows_read": ws.max_row - header_row,
        "valid": len(recipients),
        "rejected": len(rejected),
    }
    return recipients, rejected, summary
